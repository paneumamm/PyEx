"""Fill experiments 1 and 2 in the retained laboratory-report template."""

from pathlib import Path
from hashlib import sha256

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[2]
QA = ROOT / "output" / "report_qa"
TEMPLATE = Path(r"D:\code\PyExRequirement\实验报告模板.docx")
OUTPUT = ROOT / "output" / "实验报告-实验一和实验二.docx"
TEMPLATE_SHA256 = "99d12e58bfcfade7c5fb68112df940f536f281f637feeeef7f3bf334341e0261"


def write_paragraph(paragraph, content):
    paragraph.clear()
    run = paragraph.add_run(content)
    run.font.name = "宋体"
    run.font.size = Pt(11)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    paragraph.paragraph_format.space_after = Pt(4)


def answer(paragraph, content):
    write_paragraph(paragraph, "答：" + content)


def answer_png(filename, lines):
    width, line_height, padding = 1060, 42, 32
    height = len(lines) * line_height + 2 * padding
    im = Image.new("RGB", (width, height), "#17202a")
    draw = ImageDraw.Draw(im)
    font = ImageFont.truetype(r"C:\Windows\Fonts\msyh.ttc", 27)
    for index, line in enumerate(lines):
        draw.text((padding, padding + index * line_height), line,
                  fill="#eaf1f7", font=font)
    path = QA / filename
    im.save(path)
    return path


def picture(paragraph, path, width_cm):
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Cm(width_cm))


def new_after(paragraph, text, style="Normal", font_size=10.5):
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    new = type(paragraph)(element, paragraph._parent)
    new.style = style
    if text:
        run = new.add_run(text)
        run.font.name = "宋体"
        run.font.size = Pt(font_size)
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    new.paragraph_format.space_after = Pt(4)
    return new


def fill_table(table, rows):
    for row, values in zip(table.rows[1:], rows):
        for cell, content in zip(row.cells, values):
            cell.text = content
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "宋体"
                    run.font.size = Pt(10)
                    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")


def main():
    assert sha256(TEMPLATE.read_bytes()).hexdigest() == TEMPLATE_SHA256
    doc = Document(TEMPLATE)
    p = doc.paragraphs

    # Experiment 1: understanding and implementation.
    write_paragraph(p[28], "本实验把一张外卖订单从待派、配送到送达的过程变成可执行的规则，并记录实际送达时长与是否超时。骑手一次最多携带三单，因此派单必须同时维护订单状态和骑手在手订单。除自动处理样例数据外，还要让用户在图形界面中亲自派单、送达并查看实时统计。")
    write_paragraph(p[33], "程序分为三层：game_core.py 中的 Order 保存订单属性与状态，Rider 统一处理接单、送达和统计；main.py 负责读取 CSV、按 R01→R02→R03 轮流派单并生成答案；gui.py 用 tkinter 展示订单、骑手和按钮，并调用同一套核心规则。")
    write_paragraph(p[34], "自动模拟每派完 9 单结算一次，最后不足 9 单仍全部送达。实际时长按 round(18＋距离×5＋菜品数×2.5, 1) 计算；仅当实际时长严格大于承诺时长才算超时，结果写入 rider_stats.json。")
    write_paragraph(p[35], "选做的游戏模式加入 180 秒倒计时、每 3 秒自动来单、天气和骑手等级系数、在手时间修正及历史最高分存档；新一局用 game_id 避免旧定时回调干扰。")

    exp1_answers = answer_png("exp1_answers.png", [
        "===ANSWER===", "Q1=50", "Q2=19", "Q3=0.38", "Q4=R03", "Q5=0.5", "===END===",
    ])
    picture(p[40], exp1_answers, 11.5)
    picture(p[43], QA / "exp1_gui.png", 16.2)

    answer(p[48], "类属性 MAX_LOAD 表示所有 Rider 对象共享的统一容量规则，可通过 Rider.MAX_LOAD 明确读取；若放在 __init__ 中成为实例属性，每个对象会保存一份，且可能被单独改成不同上限，不利于保持规则一致。")
    answer(p[52], "@property 让 load 和 stats 像只读属性一样用 rider.load、rider.stats 访问，却仍可在读取时根据 holding、计数器实时计算。两者是随状态变化的派生值，不宜另存一份易过期的副本；普通方法则需显式加括号。")
    answer(p[56], "若容量已满却先把 order.status 改成 delivering，返回 False 后该订单仍留在待派列表，但状态显示为配送中；若先追加 holding 再判断，还会突破三单上限。应先检查容量，失败立即返回，任何对象都不修改。")
    answer(p[60], "time.sleep() 会阻塞 tkinter 的主事件循环，使窗口在等待期间不能重绘和响应点击。应使用 root.after(毫秒数, 回调函数) 安排下一次动作，让事件循环继续处理界面事件。")
    answer(p[64], "不算矛盾。main.py 使用指导书规定的固定公式，以便自动判分；gui.py 为游戏体验叠加天气、骑手等级和在手时间影响，因此同一单可能有不同实际时长。Rider.complete() 不负责计算时长，只接收调用方给出的 actual_min，并用同一条‘实际＞承诺’规则更新状态与统计。")
    fill_table(doc.tables[2], [
        ("CSV 字段 order_amount 与 Order 构造参数 amount 名称不同", "核对表头和构造函数参数，确认金额列不能直接按同名键传入", "读取时显式将 row['order_amount'] 转为 float 并传给 amount"),
        ("连续派 50 单会超过骑手三单容量", "按轮流派单检查第 10 单：R01 前三单仍在手", "每派满 9 单先让三名骑手按顺序完成在手订单，再继续派单"),
        ("重新开局后旧定时任务可能继续回调", "检查 after 回调与游戏重置后的状态关系", "每局递增 game_id，回调先核对当前局编号及 running 状态"),
    ])

    # Experiment 2: understanding and implementation.
    write_paragraph(p[72], "本实验要把 9000 条外卖订单转化为能解释运营现象的图表，而不是只输出图片。通过天气、下单时段、配送距离和商家类别观察超时率及送达时长，同时计算六项可核对的汇总结果。最后还要寻找整体规律之外的异常个体与异常日期，并说明证据和原因推测的边界。")
    write_paragraph(p[77], "用 pandas 读取 UTF-8-SIG 编码的 orders.csv，将 order_time 转为日期时间。对天气和小时用 groupby 计算 is_overtime 的均值作为超时率；对商家类别统计订单数和实际时长分布；距离与时长保留逐单数据绘制散点。")
    write_paragraph(p[78], "Matplotlib 分别输出柱状图、折线图、散点图、饼图和箱线图，统一设置中文字体、标题、坐标及单位，并在保存后关闭画布。六项答案由 len、mean、groupby 和 idxmax 等计算，超时率差换算为百分点。")
    write_paragraph(p[79], "异常分析另按骑手和日期汇总超时数、订单数与超时率：用骑手柱状图定位 R07；把每日超时率和当天天气放在同一时间轴，定位 2026-04-21。再对照其他骑手、其他晴天及送达时长，判断异常是否能由已知因素解释。")

    for index, filename, width in [
        (84, "fig1.png", 14.5), (86, "fig2.png", 14.5),
        (88, "fig3.png", 14.5), (90, "fig4.png", 11.4),
        (92, "fig5.png", 14.5),
    ]:
        picture(p[index], ROOT / "output" / "exp2" / filename, width)
    exp2_answers = answer_png("exp2_answers.png", [
        "===ANSWER===", "Q1=9000", "Q2=0.3428", "Q3=50.34",
        "Q4=正餐", "Q5=R07", "Q6=2026-04-21", "===END===",
    ])
    picture(p[95], exp2_answers, 11.5)
    picture(p[98], ROOT / "output" / "exp2" / "anomaly_rider.png", 16.0)
    picture(p[100], ROOT / "output" / "exp2" / "anomaly_daily_weather.png", 16.0)

    answer(p[105], "天气越差，超时率总体越高：晴天为 1115/3835＝29.07%，大雨为 243/306＝79.41%，大雨比晴天高 50.34 个百分点；小雨也达到 51.56%（见图 2-1）。")
    answer(p[109], "两个峰分别出现在午餐 11—12 时和晚餐 17—18 时；对应超时率约为 50.39%、52.38% 和 53.20%、50.91%，远高于 10 时的 9.03%。这些时段订单集中，商家出餐、骑手接单与道路通行可能同时承压（见图 2-2）。")
    answer(p[113], "总体呈正相关，配送距离越远，实际时长通常越长；但散点较分散，相关系数约 0.463，并非一一对应。菜品数、出餐等待、天气、骑手负载及未记录的交通条件也会影响时长（见图 2-3）。")
    answer(p[117], "箱线图除中心位置外还展示中位数、上下四分位数、须和离群点，可比较同类订单时长的波动与极端值；只画平均值的柱状图会把这些差异压成一个数。图 2-5 中正餐中位数约 34.5 分钟，高于饮品约 28.8 分钟。")
    answer(p[121], "异常骑手是 R07。先按骑手分组计算超时率，再画图 2-7 的 20 名骑手对比柱状图：R07 的 458 单中 350 单超时，达 76.42%；其他 19 人合计为 32.02%，第二高的 R15 仅 39.19%。R07 的平均距离、在手单数和承诺时长与其他人接近，可能还存在路线、取餐或班次因素，需进一步核实，不能据此认定个人责任。")
    answer(p[125], "异常日期是 2026-04-21，100 单中 91 单超时，超时率 91.00%。图 2-8 将每日超时率与天气对齐后可见，当天记录为晴天，而其余晴天订单超时率仅 27.42%；当天 20 名骑手均有超时，R07 只承接 7 单。因此它不是一般雨天波动，也不能由一名异常骑手单独解释。可能是临时出餐、调度或通行问题，具体原因需查运营记录。")
    fill_table(doc.tables[3], [
        ("order_time 原为文本，无法直接按小时/日期汇总", "检查读入后的列类型并试用 .dt 访问器", "先 pd.to_datetime，再分别使用 .dt.hour 和 .dt.date 分组"),
        ("超时率差容易把比例与百分点混淆", "核对 Q3 单位：两组 mean 得到 0—1 比例", "先用大雨率减晴天率，再乘 100，保留两位小数，得 50.34 个百分点"),
        ("单看每日超时峰值无法判断天气能否解释", "把峰值日天气与其他晴天逐单超时率对比", "绘制超时率—天气双层时间图，并核查 4 月 21 日 91/100 单及全员超时记录"),
    ])

    # Retain captions; add optional-task evidence directly after the GUI caption.
    optional_heading = new_after(p[44], "③ 选做功能", "Heading 3", 11)
    optional_image = new_after(optional_heading, "", "截图位")
    picture(optional_image, QA / "exp1_optional.png", 16.2)
    optional_caption = new_after(optional_image, "图 1-3　选做功能", "Caption", 10.5)
    optional_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    new_after(optional_caption, "界面顶部展示倒计时、天气、总超时率和得分；此图为送达一单后的状态。程序另实现自动来单、骑手等级影响和历史最佳成绩存档。")

    new_after(p[99], "R07 超时率 76.42%，比其余骑手合计高 44.40 个百分点；图中红色柱与虚线基准突出这一差距。")
    new_after(p[101], "4 月 21 日晴天却达到 91.00% 超时率，比其余晴天订单高 63.58 个百分点；图中将峰值与天气放在同一时间轴核对。")

    # The template's empty handwriting lines are no longer needed after answers are filled.
    for index in (49, 50, 53, 54, 57, 58, 61, 62, 65, 66,
                  106, 107, 110, 111, 114, 115, 118, 119, 122, 123, 126, 127):
        element = p[index]._element
        element.getparent().remove(element)
    for index in (48, 52, 56, 60, 64, 105, 109, 113, 117, 121, 125):
        p[index].paragraph_format.keep_together = True

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)
    print("bytes", OUTPUT.stat().st_size)
    print("source_sha256", sha256(TEMPLATE.read_bytes()).hexdigest())


if __name__ == "__main__":
    main()
